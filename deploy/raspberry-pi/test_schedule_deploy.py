import importlib.util
from datetime import datetime
from pathlib import Path
import json
import subprocess
import tempfile
import unittest
from unittest.mock import Mock, patch

from openpyxl import Workbook

spec = importlib.util.spec_from_file_location('schedule_deploy', Path(__file__).with_name('schedule_deploy.py'))
schedule = importlib.util.module_from_spec(spec)
spec.loader.exec_module(schedule)


class ScheduleDeploymentTests(unittest.TestCase):
    def fixture(self, tmp):
        home = Path(tmp)
        root, checkout = home/'live', home/'checkout'
        for base in [root, checkout]:
            (base/'organizer').mkdir(parents=True)
            (base/'organizer/web.py').write_text('existing code')
        private = home/'.config/project-hub'
        private.mkdir(parents=True)
        (private/'library.env').write_text('PRIVATE_SETTING=preserved\n')
        (root/'token.json').write_text('bot token preserved')
        (root/'project_hub.db').write_bytes(b'live database')
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = '일별 시간표'
        sheet.cell(7, 2, datetime(2027, 2, 1))
        sheet.cell(7, 4, 1)
        sheet.cell(7, 5, '09:10~10:00')
        sheet.cell(7, 6, '샘플 과목')
        source = home/'source.xlsx'
        workbook.save(source)
        workbook.close()
        return home, root, checkout, source.read_bytes()

    def test_drive_download_uses_exact_file_and_existing_read_auth(self):
        service = Mock()
        files = service.files.return_value
        files.get.return_value.execute.return_value = dict(name='일정표.xlsx', mimeType='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', size='3')
        files.get_media.return_value.execute.return_value = b'zip'
        with patch.object(schedule, 'photo_service', return_value=service) as auth:
            self.assertEqual(schedule.download_source('chosen-file', Path('/private/token')), ('일정표.xlsx', b'zip'))
        auth.assert_called_once_with(Path('/private/token'))
        self.assertEqual(files.get.call_args.kwargs['fileId'], 'chosen-file')
        self.assertEqual(files.get_media.call_args.kwargs['fileId'], 'chosen-file')
        files.get.return_value.execute.return_value['size'] = str(9*1024*1024)
        files.get_media.reset_mock()
        with patch.object(schedule, 'photo_service', return_value=service), self.assertRaises(schedule.SetupError):
            schedule.download_source('chosen-file', Path('/private/token'))
        files.get_media.assert_not_called()

    def test_invalid_excel_does_not_publish_or_deploy(self):
        with tempfile.TemporaryDirectory() as tmp:
            home, root, checkout, data = self.fixture(tmp)
            with patch.object(schedule, 'download_source', return_value=('sample.xlsx', b'invalid')), \
                 patch.object(schedule.subprocess, 'run') as deploy:
                with self.assertRaises(Exception):
                    schedule.prepare_and_apply(checkout, root, home, 'chosen-file')
            deploy.assert_not_called()
            self.assertFalse((home/'.config/project-hub/schedule/main.json').exists())

    def test_download_apply_and_service_readback_preserve_other_data(self):
        with tempfile.TemporaryDirectory() as tmp:
            home, root, checkout, data = self.fixture(tmp)
            directory = home/'.config/project-hub/schedule'
            response = Mock()
            response.__enter__ = Mock(return_value=response)
            response.__exit__ = Mock(return_value=False)
            source_path = home/'source.xlsx'
            converted = schedule.parse_schedule(source_path)
            response.read.return_value = json.dumps({'available': True, 'source': 'Drive schedule.xlsx', 'days': converted['days']}).encode()
            with patch.object(schedule, 'download_source', return_value=('Drive schedule.xlsx', data)), \
                 patch.object(schedule.subprocess, 'run') as deploy, \
                 patch.object(schedule.urllib.request, 'urlopen', return_value=response):
                schedule.prepare_and_apply(checkout, root, home, 'chosen-file')
            deploy.assert_called_once()
            self.assertIn('--apply', deploy.call_args.args[0])
            manifest = json.loads((directory/'main.json').read_text())
            self.assertEqual(manifest['file_id'], 'chosen-file')
            self.assertEqual((directory/(manifest['sha256']+'.xlsx')).read_bytes(), data)
            self.assertEqual((directory/'main.json').stat().st_mode & 0o777, 0o600)
            self.assertEqual((root/'token.json').read_text(), 'bot token preserved')
            self.assertEqual((root/'project_hub.db').read_bytes(), b'live database')
            self.assertEqual((home/'.config/project-hub/library.env').read_text(), 'PRIVATE_SETTING=preserved\n')

    def test_apply_failure_restores_prior_pointer(self):
        for existing in [False, True]:
            with self.subTest(existing=existing), tempfile.TemporaryDirectory() as tmp:
                home, root, checkout, data = self.fixture(tmp)
                directory = home/'.config/project-hub/schedule'
                directory.mkdir()
                manifest = directory/'main.json'
                old = b'{"previous":"configuration"}'
                if existing:
                    manifest.write_bytes(old)
                with patch.object(schedule, 'download_source', return_value=('sample.xlsx', data)), \
                     patch.object(schedule.subprocess, 'run', side_effect=subprocess.CalledProcessError(1, 'deploy')):
                    with self.assertRaises(subprocess.CalledProcessError):
                        schedule.prepare_and_apply(checkout, root, home, 'chosen-file')
                if existing:
                    self.assertEqual(manifest.read_bytes(), old)
                else:
                    self.assertFalse(manifest.exists())
