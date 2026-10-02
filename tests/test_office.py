from contextlib import closing
from datetime import datetime
import io
from pathlib import Path
import unittest
from unittest.mock import patch
import zipfile

from docx import Document
from openpyxl import Workbook
import test_content as fixtures
import test_file_search as search_fixtures
from bot import database, content_search as content, content_extract as extraction, office_extract as office


def docx_bytes():
    document = Document()
    document.add_paragraph('첫 번째 문단')
    document.add_paragraph('서보 모터 문단')
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = 'PLC 입력'
    table.cell(0, 1).text = '스캔 주기'
    table.cell(1, 0).merge(table.cell(1, 1)).text = '병합 셀 내용'
    table.cell(0, 1).add_table(rows=1, cols=1).cell(0, 0).text = '중첩 표 설명'
    document.sections[0].header.paragraphs[0].text = '본문 아닌 머리말'
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def xlsx_bytes():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = '설비 설정'
    sheet['A1'] = '서보 속도'
    sheet['B2'] = 1200
    sheet['C3'] = '=SUM(B2,1)'
    sheet['D4'] = datetime(2026, 10, 2)
    sheet['E5'] = True
    hidden = workbook.create_sheet('숨김 자료')
    hidden.sheet_state = 'hidden'
    hidden['A1'] = '숨김 비밀'
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def rewrite_zip(data, name, change):
    result = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as source, zipfile.ZipFile(result, 'w', zipfile.ZIP_DEFLATED) as target:
        for item in source.infolist():
            value = source.read(item.filename)
            target.writestr(item, change(value) if item.filename == name else value)
    return result.getvalue()


class OfficeExtractionTests(fixtures.ContentFixture, unittest.TestCase):
    def test_word_paragraph_table_nested_table_and_merged_cell(self):
        result = self.index(self.file(docx_bytes(), '설명.DOCX'))
        self.assertEqual(result['status'], 'indexed')
        rows, total, _ = content.find_content(1, [10], '서보')
        self.assertEqual((total, rows[0]['location']), (1, 'Word 문단 2'))
        self.assertIn('Word 표 1 · 1행 1열', content.find_content(1, [10], 'PLC')[0][0]['location'])
        self.assertIn('표 1 · 1행 2열 표 1', content.find_content(1, [10], '중첩')[0][0]['location'])
        self.assertEqual(sum('병합 셀 내용' in body for _, body in result['pages']), 1)
        self.assertEqual(content.find_content(1, [10], '머리말')[1], 0)

    def test_excel_sheet_cell_formula_dates_and_hidden_sheet(self):
        self.index(self.file(xlsx_bytes(), '설비.xlsx'))
        rows, total, _ = content.find_content(1, [10], '서보')
        self.assertEqual((total, rows[0]['location']), (1, 'Excel 설비 설정 · A1'))
        rows, _, _ = content.find_content(1, [10], 'SUM')
        self.assertIn('C3 · 수식', rows[0]['location'])
        self.assertIn('=SUM(B2,1)', rows[0]['excerpt'])
        self.assertIn('D4', content.find_content(1, [10], '2026-10-02')[0][0]['location'])
        self.assertIn('B2', content.find_content(1, [10], '1200')[0][0]['location'])
        self.assertIn('E5', content.find_content(1, [10], 'True')[0][0]['location'])
        self.assertEqual(content.find_content(1, [10], '숨김 비밀')[1], 0)

    def test_wrong_sheet_dimensions_do_not_hide_data(self):
        data = rewrite_zip(xlsx_bytes(), 'xl/worksheets/sheet1.xml',
                           lambda value: value.replace(b'ref="A1:E5"', b'ref="A1:A1"'))
        self.index(self.file(data, 'wrong-dimensions.xlsx'))
        self.assertEqual(content.find_content(1, [10], 'True')[1], 1)

    def test_excel_and_word_limits_report_partial(self):
        row = self.file(xlsx_bytes(), 'large.xlsx')
        with patch.object(office, 'MAX_EXCEL_ROWS', 2), patch.object(office, 'MAX_EXCEL_COLS', 2):
            result = office.extract_office(Path(row['local_path']).read_bytes(), '.xlsx', extraction.clean_text, extraction.MAX_CHARS)
        self.assertEqual(result['status'], 'partial')
        self.assertEqual(len(result['pages']), 2)
        with patch.object(office, 'MAX_PARTS', 1):
            result = office.extract_office(docx_bytes(), '.docx', extraction.clean_text, extraction.MAX_CHARS)
        self.assertEqual(result['status'], 'partial')
        self.assertEqual(len(result['pages']), 1)
        result = office.extract_office(docx_bytes(), '.docx', extraction.clean_text, 3)
        self.assertEqual(result['status'], 'partial')
        self.assertLessEqual(sum(len(text) for _, text in result['pages']), 3)

    def test_zip_expansion_and_encrypted_package_are_excluded(self):
        with patch.object(office, 'MAX_UNPACKED_BYTES', 100):
            self.assertEqual(office.extract_office(docx_bytes(), '.docx', extraction.clean_text, 100)['status'], 'too_large')
        result = office.extract_office(b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1' + b'test', '.xlsx', extraction.clean_text, 100)
        self.assertEqual(result['status'], 'encrypted')

    def test_incorrect_dimensions_still_report_column_truncation(self):
        data = rewrite_zip(xlsx_bytes(), 'xl/worksheets/sheet1.xml',
                           lambda value: value.replace(b'ref="A1:E5"', b'ref="A1:A1"'))
        with patch.object(office, 'MAX_EXCEL_COLS', 2):
            result = office.extract_office(data, '.xlsx', extraction.clean_text, extraction.MAX_CHARS)
        self.assertEqual(result['status'], 'partial')
        self.assertNotIn('=SUM', str(result['pages']))

    def test_word_external_entities_cannot_read_local_files(self):
        private = self.storage.parent / 'private-entity.txt'
        private.write_text('ENTITY_PRIVATE_CONTENT', encoding='utf-8')
        def entity(value):
            head, body = value.split(b'?>', 1)
            declaration = ('<!DOCTYPE w:document [<!ENTITY secret SYSTEM "' + private.as_uri() + '">]>').encode()
            return head + b'?>' + declaration + body.replace('첫 번째 문단'.encode(), b'&secret;')
        data = rewrite_zip(docx_bytes(), 'word/document.xml', entity)
        result = extraction.extract(self.job(self.file(data, 'entity.docx')))
        self.assertNotIn('ENTITY_PRIVATE_CONTENT', str(result))

    def test_legacy_and_macro_extensions_stay_unsupported(self):
        for extension in ('doc', 'xls', 'docm', 'xlsm'):
            self.assertEqual(self.index(self.file(b'old file', 'file.' + extension))['status'], 'unsupported')


class OfficeIndexTests(fixtures.ContentFixture, unittest.TestCase):
    def test_only_newly_supported_office_files_need_reindexing(self):
        text = self.file(b'old TXT remains searchable')
        self.index(text)
        word = self.file(docx_bytes(), 'old.docx')
        # Simulate the previous release having marked this file unsupported.
        with closing(database.get_connection()) as conn, conn:
            conn.execute('INSERT INTO content_documents (file_id,source_sha256,index_version,status) VALUES (?,?,1,?)',
                         (word['id'], word['sha256'], 'unsupported'))
        self.assertEqual([r['id'] for r in content.due_files()], [word['id']])
        self.assertEqual(content.find_content(1, [10], 'TXT')[1], 1)
        self.index(word)
        self.assertEqual(content.due_files(), [])
        self.assertEqual(content.find_content(1, [10], '서보')[1], 1)

    def test_locations_and_counts_are_scoped_to_visible_guild_channels(self):
        self.index(self.file(docx_bytes(), 'allowed.docx'))
        self.index(self.file(xlsx_bytes(), 'private.xlsx', channel='20'))
        self.index(self.file(xlsx_bytes(), 'foreign.xlsx', guild='2'))
        rows, total, states = content.find_content(1, [10], '서보')
        self.assertEqual((total, states), (1, {'indexed': 1}))
        self.assertEqual(rows[0]['original_filename'], 'allowed.docx')
        self.assertNotIn('Excel', str(rows))

    def test_schema_migration_keeps_old_content_pages_insert_compatible(self):
        with closing(database.get_connection()) as conn, conn:
            conn.execute('DROP TABLE content_locations')
        content.initialize()
        with closing(database.get_connection()) as conn, conn:
            conn.execute('INSERT INTO content_pages VALUES (999,1,?,?)', ('legacy', 'legacy'))
            self.assertEqual(conn.execute('SELECT COUNT(*) FROM content_pages WHERE file_id=999').fetchone()[0], 1)


class OfficeWorkerTests(fixtures.ContentFixture, unittest.IsolatedAsyncioTestCase):
    async def test_actual_isolated_workers_extract_word_and_excel(self):
        self.file(docx_bytes(), 'manual.docx')
        self.file(xlsx_bytes(), 'control.xlsx')
        self.assertEqual(await content.ContentIndexer(self.storage).scan_once(), 2)
        rows, total, _ = content.find_content(1, [10], '서보')
        self.assertEqual(total, 2)
        self.assertEqual({r['location'] for r in rows}, {'Word 문단 2', 'Excel 설비 설정 · A1'})
        self.assertEqual(await content.ContentIndexer(self.storage).scan_once(), 0)

    async def test_malformed_office_file_does_not_block_other_documents(self):
        self.file(b'broken zip', 'broken.docx')
        self.file(b'servo success', 'normal.txt')
        self.assertEqual(await content.ContentIndexer(self.storage).scan_once(), 2)
        self.assertEqual(content.find_content(1, [10], 'servo')[1], 1)


class OfficeCommandTests(fixtures.ContentFixture, unittest.IsolatedAsyncioTestCase):
    asyncSetUp = search_fixtures.CommandTests.asyncSetUp
    interaction = search_fixtures.CommandTests.interaction

    async def test_existing_command_displays_word_and_excel_locations(self):
        self.index(self.file(docx_bytes(), 'manual.docx'))
        self.index(self.file(xlsx_bytes(), 'control.xlsx'))
        interaction = self.interaction()
        with patch.object(self.client, 'get_channel', return_value=search_fixtures.channel()):
            await self.commands.tree.get_command('내용검색').callback(interaction, '서보')
        embed = interaction.edit_original_response.call_args.kwargs['embed'].to_dict()
        interaction.response.defer.assert_awaited_once_with(ephemeral=True, thinking=True)
        self.assertIn('Word 문단 2', str(embed))
        self.assertIn('Excel 설비 설정 · A1', str(embed))
        self.assertNotIn('TXT 본문', str(embed))
        self.assertTrue(all(len(field['value']) <= 1024 for field in embed['fields']))
