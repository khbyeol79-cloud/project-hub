from datetime import datetime, timezone
from pathlib import Path
import sqlite3
import tempfile
import unittest

from openpyxl import Workbook

from organizer.schedule import parse_schedule, schedule_snapshot
from organizer.web import create_app


def workbook_fixture(path):
    workbook = Workbook()
    overview = workbook.active
    overview.title = '월별 개요'
    overview.cell(7, 2, 2)
    overview.cell(7, 4, '휴강')
    overview.cell(7, 5, '수료')
    sheet = workbook.create_sheet('세부일정 및 일별 시간표')
    sheet.cell(3, 1, '훈련과정 : 샘플 과정')
    sheet.cell(5, 1, '훈련기간 : 2027.02.01~2027.02.03')
    times = ['09:10~10:00', '10:10~11:00', '11:10~12:00', '12:10~13:00',
             '14:00~14:50', '15:00~15:50', '16:00~16:50', '17:00~17:50']
    for day in range(1, 4):
        row = 7 + (day - 1) * 8
        sheet.cell(row, 2, datetime(2027, 2, day))
        for period in range(1, 9):
            r = row + period - 1
            sheet.cell(r, 4, period)
            sheet.cell(r, 5, times[period - 1])
            if day == 2 or day == 3 and period > 5 or day == 1 and period == 1:
                continue
            sheet.cell(r, 6, '과목 A' if period < 8 else '과목 B')
            sheet.cell(r, 8, '공개하지 않을 강사')
            sheet.cell(r, 9, '강의실 A')
    workbook.save(path)
    workbook.close()


class ScheduleTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        self.storage = self.root/'storage'
        self.storage.mkdir()
        self.path = self.storage/'sample.xlsx'
        workbook_fixture(self.path)

    def snapshot(self, now):
        return schedule_snapshot(self.path, '샘플 일정표.xlsx', now)

    def test_window_moves_at_korean_monday_midnight(self):
        sunday = self.snapshot(datetime(2026, 10, 11, 14, 59, tzinfo=timezone.utc))
        monday = self.snapshot(datetime(2026, 10, 11, 15, 0, tzinfo=timezone.utc))
        self.assertEqual((sunday['today'], sunday['week_start'], sunday['week_end']),
                         ('2026-10-11', '2026-10-05', '2026-10-25'))
        self.assertEqual((monday['today'], monday['week_start'], monday['week_end']),
                         ('2026-10-12', '2026-10-12', '2026-11-01'))

    def test_window_keeps_rolling_across_year(self):
        data = self.snapshot(datetime(2026, 12, 31, tzinfo=timezone.utc))
        self.assertEqual((data['week_start'], data['week_end']), ('2026-12-28', '2027-01-17'))
        self.assertNotIn('2027-02-04', data['days'])

    def test_parse_blanks_holidays_lunch_and_mixed_last_period(self):
        data = parse_schedule(self.path)
        self.assertEqual((data['teaching_days'], data['teaching_periods']), (2, 12))
        self.assertEqual(data['days']['2027-02-02'], {'note': '휴강', 'lessons': []})
        mixed = data['days']['2027-02-01']['lessons']
        self.assertEqual([(b['subject'], b['start_period'], b['end_period']) for b in mixed],
                         [('과목 A', 2, 4), ('과목 A', 5, 7), ('과목 B', 8, 8)])
        self.assertEqual(mixed[0]['start'], '10:10')
        self.assertEqual(mixed[-1]['start'], '17:00')
        self.assertEqual(data['days']['2027-02-03']['lessons'][-1]['end'], '14:50')
        self.assertNotIn('공개하지 않을 강사', str(data))

    def test_cache_is_read_only_and_invalidated_when_original_changes(self):
        data = self.snapshot(datetime.now(timezone.utc))
        data['days']['2027-02-01']['lessons'].clear()
        self.assertTrue(self.snapshot(datetime.now(timezone.utc))['days']['2027-02-01']['lessons'])
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = '일별 시간표'
        sheet.cell(7, 2, datetime(2027, 3, 1))
        sheet.cell(7, 4, 1)
        sheet.cell(7, 5, '09:10~10:00')
        sheet.cell(7, 6, '변경 과목')
        workbook.save(self.path)
        workbook.close()
        self.assertEqual(self.snapshot(datetime.now(timezone.utc))['start'], '2027-03-01')

    def test_endpoint_access_scope_latest_source_and_path_boundary(self):
        db = self.root/'source.db'
        with sqlite3.connect(db) as connection:
            connection.executescript('''CREATE TABLE files(id INTEGER PRIMARY KEY, original_filename TEXT,
              discord_channel TEXT, discord_channel_id TEXT, category TEXT, uploaded_at TEXT,
              file_size_bytes INTEGER, version INTEGER, duplicate_type TEXT, sha256 TEXT, local_path TEXT);
              CREATE TABLE content_documents(file_id INTEGER, source_sha256 TEXT, status TEXT);''')
            for id, name, uploaded in [(1, '샘플 일정표.xlsx', '2027-01-01'), (2, '수정 일정표.xlsx', '2027-01-02')]:
                connection.execute('INSERT INTO files VALUES(?,?,?,?,?,?,?,?,?,?,?)',
                    (id, name, 'test', '1', 'plc', uploaded, 100, 1, None, 'test-hash', str(self.path)))
        app = create_app({'TESTING': True, 'DB_PATH': db, 'STORAGE': self.storage, 'SECRET_KEY': 'test-only'})
        client = app.test_client()
        self.assertEqual(client.get('/library/api/schedule').status_code, 401)
        with client.get('/library/assets/schedule.js') as asset:
            self.assertEqual(asset.status_code, 200)
        app.config['PUBLIC_ACCESS'] = True
        response = client.get('/library/api/schedule')
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json['source'], '수정 일정표.xlsx')
        self.assertEqual(response.headers['Cache-Control'], 'no-store')
        for team in ['1a', '1b']:
            data = client.get('/library/api/schedule?project=additional&team='+team).json
            self.assertFalse(data['available'])
            self.assertNotIn('days', data)
        self.assertEqual(client.get('/library/api/schedule?project=invalid').status_code, 400)
        self.assertEqual(client.get('/library/api/schedule?project=additional&team=wrong').status_code, 400)
        with sqlite3.connect(db) as connection:
            connection.execute('UPDATE files SET local_path=? WHERE id=2', (str(self.root/'outside.xlsx'),))
        self.assertEqual(client.get('/library/api/schedule').status_code, 404)
