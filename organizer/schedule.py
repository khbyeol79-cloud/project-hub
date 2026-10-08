"""Render a collected education workbook without publishing its contents as code."""
from copy import deepcopy
from datetime import date, datetime, timedelta
from functools import lru_cache
from pathlib import Path
import re
from zipfile import ZipFile
from zoneinfo import ZoneInfo

import openpyxl


class ScheduleError(ValueError):
    pass


def parse_schedule(path):
    path = Path(path)
    if path.stat().st_size > 8 * 1024 * 1024:
        raise ScheduleError('일정표 파일이 너무 큽니다.')
    with ZipFile(path) as archive:
        if sum(entry.file_size for entry in archive.infolist()) > 32 * 1024 * 1024:
            raise ScheduleError('일정표 압축 해제 크기가 너무 큽니다.')
    workbook = openpyxl.load_workbook(path, data_only=True, read_only=True, keep_links=False)
    try:
        sheet = next((s for s in workbook if '일별' in s.title and '시간표' in s.title), None)
        if sheet is None or sheet.max_row > 10000 or sheet.max_column > 512:
            raise ScheduleError('일별 시간표 시트를 확인해 주세요.')
        course, start, end, days = '', None, None, {}
        current = None
        for cells in sheet.iter_rows(max_col=9, values_only=True):
            heading = str(cells[0] or '').strip()
            if heading.startswith('훈련과정'):
                course = heading.split(':', 1)[-1].strip()
            if heading.startswith('훈련기간'):
                dates = re.findall(r'\d{4}[.-]\d{2}[.-]\d{2}', heading)
                if len(dates) == 2:
                    start, end = (date.fromisoformat(v.replace('.', '-')) for v in dates)
            if isinstance(cells[1], (datetime, date)):
                current = cells[1].date() if isinstance(cells[1], datetime) else cells[1]
                days.setdefault(current.isoformat(), {'note': '', 'lessons': []})
            if current is None or not cells[5]:
                continue
            period = cells[3]
            times = re.fullmatch(r'(\d{2}:\d{2})\s*[~–-]\s*(\d{2}:\d{2})', str(cells[4] or '').strip())
            if not isinstance(period, (int, float)) or int(period) != period or not 1 <= period <= 8 or not times:
                raise ScheduleError('교시와 시간을 확인해 주세요.')
            period = int(period)
            subject, room = str(cells[5]).strip(), str(cells[8] or '').strip()
            lessons = days[current.isoformat()]['lessons']
            if lessons and lessons[-1]['subject'] == subject and lessons[-1]['room'] == room and lessons[-1]['end_period'] == period - 1 and period != 5:
                lessons[-1].update(end_period=period, end=times[2])
            else:
                lessons.append(dict(start_period=period, end_period=period, start=times[1], end=times[2], subject=subject, room=room))
        teaching = [date.fromisoformat(key) for key, day in days.items() if day['lessons']]
        if not teaching:
            raise ScheduleError('일정표에 등록된 수업이 없습니다.')
        start, end = start or min(teaching), end or max(teaching)
        if start > end or (end - start).days > 730:
            raise ScheduleError('훈련 기간을 확인해 주세요.')
        days = {key: day for key, day in days.items() if start.isoformat() <= key <= end.isoformat()}
        # The compact monthly overview supplies admission, holidays and closure notes.
        for overview in workbook:
            if overview is sheet or overview.max_row > 100 or overview.max_column > 512:
                continue
            for cells in overview.iter_rows(max_col=33, values_only=True):
                month = cells[1]
                if not isinstance(month, int) or not 1 <= month <= 12:
                    continue
                year = start.year + (1 if end.year > start.year and month < start.month else 0)
                for number, value in enumerate(cells[2:], 1):
                    if not isinstance(value, str) or not value.strip():
                        continue
                    try:
                        key = date(year, month, number).isoformat()
                    except ValueError:
                        continue
                    if start.isoformat() <= key <= end.isoformat():
                        text = value.strip()
                        day = days.setdefault(key, {'note': '', 'lessons': []})
                        day['note'] = {'입학': '입학식', '대체': '대체휴일'}.get(text, text)
        return dict(course=course or '교육 일정', start=start.isoformat(), end=end.isoformat(), days=days,
                    teaching_days=sum(bool(day['lessons']) for day in days.values()),
                    teaching_periods=sum(b['end_period']-b['start_period']+1 for day in days.values() for b in day['lessons']))
    finally:
        workbook.close()


@lru_cache(maxsize=8)
def _cached_schedule(path, modified, size):
    return parse_schedule(path)


def schedule_snapshot(path, source, now=None):
    path = Path(path)
    stat = path.stat()
    data = deepcopy(_cached_schedule(str(path.resolve()), stat.st_mtime_ns, stat.st_size))
    now = now or datetime.now(ZoneInfo('Asia/Seoul'))
    today = now.astimezone(ZoneInfo('Asia/Seoul')).date()
    monday = today - timedelta(days=today.weekday())
    data.update(source=source, available=True, timezone='Asia/Seoul', today=today.isoformat(),
                week_start=monday.isoformat(), week_end=(monday+timedelta(days=20)).isoformat())
    return data
