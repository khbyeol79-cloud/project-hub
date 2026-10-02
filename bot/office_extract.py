"""Read DOCX/XLSX content without executing formulas, macros, or external links."""
import io
import zipfile

MAX_PARTS = 2000
MAX_WORD_BLOCKS = 4000
MAX_EXCEL_ROWS = 5000
MAX_EXCEL_COLS = 256
MAX_SHEETS = 30
MAX_UNPACKED_BYTES = 80 * 1024 * 1024


def word_parts(container, prefix='Word'):
    from docx.text.paragraph import Paragraph
    paragraph_number, table_number = 0, 0
    for block in container.iter_inner_content():
        if isinstance(block, Paragraph):
            paragraph_number += 1
            yield f'{prefix} 문단 {paragraph_number}', block.text
        else:
            table_number += 1
            seen_cells = set()
            for row_number, row in enumerate(block.rows, 1):
                for column_number, cell in enumerate(row.cells, 1):
                    if cell._tc in seen_cells:
                        continue
                    seen_cells.add(cell._tc)
                    label = f'{prefix} 표 {table_number} · {row_number}행 {column_number}열'
                    yield from word_parts(cell, label)


def extract_office(data, suffix, clean_text, max_chars):
    # Password-protected modern Office files use an OLE encrypted package.
    if data.startswith(b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1'):
        return {'status': 'encrypted', 'pages': []}
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        entries = archive.infolist()
        if (len(entries) > 5000 or sum(e.file_size for e in entries) > MAX_UNPACKED_BYTES
                or any(e.file_size > 32 * 1024 * 1024 and e.filename.endswith('.xml') for e in entries)):
            return {'status': 'too_large', 'pages': []}
        if any(e.flag_bits & 1 for e in entries):
            return {'status': 'encrypted', 'pages': []}
    pages, locations, used = [], {}, 0
    partial = False

    def add(label, value):
        nonlocal used, partial
        text = clean_text(str(value))
        if not text:
            return True
        if len(pages) >= MAX_PARTS or used >= max_chars:
            partial = True
            return False
        limit = min(100_000, max_chars - used)
        if len(text) > limit:
            partial = True
        text = text[:limit]
        number = len(pages) + 1
        pages.append([number, text])
        locations[str(number)] = clean_text(label)[:240]
        used += len(text)
        return True

    if suffix == '.docx':
        from docx import Document
        document = Document(io.BytesIO(data))
        for count, (label, value) in enumerate(word_parts(document), 1):
            if count > MAX_WORD_BLOCKS:
                partial = True
                break
            if not add(label, value):
                break
    else:
        from openpyxl import load_workbook
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
        try:
            visible = [sheet for sheet in workbook.worksheets if sheet.sheet_state == 'visible']
            partial = len(visible) > MAX_SHEETS
            stopped = False
            for sheet in visible[:MAX_SHEETS]:
                if (sheet.max_row or 0) > MAX_EXCEL_ROWS or (sheet.max_column or 0) > MAX_EXCEL_COLS:
                    partial = True
                # Producers sometimes write an incorrect A1:A1 dimension.
                sheet.reset_dimensions()
                for row_number, row in enumerate(sheet.iter_rows(), 1):
                    if row_number > MAX_EXCEL_ROWS:
                        partial = True
                        break
                    if len(row) > MAX_EXCEL_COLS:
                        partial = True
                    for cell in row[:MAX_EXCEL_COLS]:
                        if cell.value is None:
                            continue
                        value = cell.value
                        # Non-string array/data-table formula objects have no searchable value.
                        if not isinstance(value, (str, int, float, bool)):
                            if hasattr(value, 'isoformat'):
                                value = value.isoformat()
                            else:
                                partial = True
                                continue
                        label = f'Excel {sheet.title} · {cell.coordinate}'
                        if cell.data_type == 'f':
                            label += ' · 수식'
                        if not add(label, value):
                            stopped = True
                            break
                    if stopped:
                        break
                if stopped:
                    break
        finally:
            workbook.close()
    return {'status': 'partial' if partial else ('indexed' if pages else 'no_text'),
            'pages': pages, 'locations': locations}
